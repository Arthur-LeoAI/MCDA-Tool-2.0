"""Arbitrary tables (/api/v1/tables/*): inspect an upload, then build a project from column choices.

See docs/design/s2-data-preparation.md. Rows and columns are 0-based here.
"""

from __future__ import annotations

import datetime as dt
import io
import math
from typing import Annotated, Literal

import openpyxl
from fastapi import APIRouter, File, Query, Request, UploadFile
from pydantic import BaseModel, Field

from mcda_core import McdaError, ProblemValidationError
from mcda_core.io import looks_like_legacy, parse_legacy_rows, parse_rows, read_csv_rows, read_xlsx_rows
from mcda_core.model import Direction
from mcda_core.project import Project
from mcda_core.tables import ColumnSpec, build_project, describe_columns, guess_header_row, guess_name_column

from .errors import ApiError
from .projects import project_from_upload
from .routes import ERRORS, Upload, compute, receive_upload, unreadable
from .schemas import ErrorResponse

router = APIRouter(prefix="/api/v1/tables", tags=["tables"], responses=ERRORS)

MAX_BODY_ROWS = 500
MAX_COLUMNS = 200

Cell = bool | float | str | None


class TableColumn(BaseModel):
    index: int
    header: str
    detected_type: Literal["numeric", "ordinal", "categorical", "boolean", "text", "empty"]
    present: int
    numeric_share: float
    percent: bool
    categories: list[str]
    examples: list[str]


class TableInspection(BaseModel):
    filename: str
    format: Literal["table", "template", "legacy"]
    project: Project | None = None
    """Set for templates and legacy workbooks, which need no column choices."""
    sheets: list[str] = []
    sheet: str | None = None
    header_row: int = 0
    name_column: int | None = None
    """Suggested column with the alternatives' names."""
    columns: list[TableColumn] = []
    rows: list[list[Cell]] = []
    """All rows from the top, up to 500 rows below the header row."""
    truncated: bool = False


class TableColumnSpec(BaseModel):
    index: Annotated[int, Field(ge=0)]
    type: Literal["numeric", "ordinal", "categorical", "boolean"]
    name: Annotated[str, Field(max_length=200)] | None = None
    direction: Direction = Direction.BENEFIT
    categories: Annotated[list[Annotated[str, Field(max_length=500)]], Field(max_length=100)] | None = None
    """Ordinal: worst first."""


class TableBuildRequest(BaseModel):
    rows: Annotated[list[Annotated[list[Cell], Field(max_length=MAX_COLUMNS)]], Field(max_length=MAX_BODY_ROWS + 20)]
    header_row: Annotated[int, Field(ge=0)]
    name_column: Annotated[int, Field(ge=0)]
    columns: Annotated[list[TableColumnSpec], Field(min_length=1, max_length=MAX_COLUMNS)]
    filename: str | None = None
    sheet: str | None = None


def _json_cell(value: object) -> Cell:
    if value is None or isinstance(value, (bool, str)):
        return value
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(value) else None
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    return str(value)


@router.post(
    "/inspect",
    response_model=TableInspection,
    responses={413: {"model": ErrorResponse}, 415: {"model": ErrorResponse}},
)
def inspect_table(
    http_request: Request,
    file: Annotated[UploadFile, File()],
    sheet: Annotated[str | None, Query()] = None,
    header_row: Annotated[int | None, Query(ge=0)] = None,
) -> TableInspection:
    """Read any spreadsheet. Templates and legacy workbooks come back as projects;
    other tables come back with detected structure for the user to confirm."""
    received = receive_upload(http_request, file)
    sheets: list[str] = []
    try:
        if received.suffix == ".csv":
            rows = read_csv_rows(received.data)
        else:
            workbook = openpyxl.load_workbook(io.BytesIO(received.data), read_only=True)
            sheets = list(workbook.sheetnames)
            workbook.close()
            if sheet is not None and sheet not in sheets:
                raise McdaError("sheet_not_found", sheet=sheet)
            rows = read_xlsx_rows(io.BytesIO(received.data), prefer_sheet=sheet)
    except (McdaError, ProblemValidationError):
        raise
    except Exception as exc:  # corrupt or non-spreadsheet content
        raise unreadable(received, exc) from exc

    # Recognised layouts need no column choices (only when the user did not pick a sheet).
    first_label = str(rows[0][0]).strip().lower() if rows and rows[0] and rows[0][0] is not None else ""
    if sheet is None and header_row is None and (looks_like_legacy(rows) or first_label == "criterion"):
        legacy = parse_legacy_rows(rows) if looks_like_legacy(rows) else None
        upload = Upload(received.filename, legacy.problem if legacy else parse_rows(rows), legacy)
        source_format, project = project_from_upload(upload)
        return TableInspection(filename=received.filename, format=source_format, project=project, sheets=sheets)

    if not any(any(v is not None and str(v).strip() for v in row) for row in rows):
        raise ApiError(422, "table_empty")
    header = guess_header_row(rows) if header_row is None else header_row
    if header >= len(rows):
        raise McdaError("table_header_row_invalid", row=header + 1)
    limit = header + 1 + MAX_BODY_ROWS
    kept = [[_json_cell(v) for v in row[:MAX_COLUMNS]] for row in rows[:limit]]
    columns = describe_columns(kept, header)
    body_rows = len([r for r in kept[header + 1 :] if any(v is not None for v in r)])
    return TableInspection(
        filename=received.filename,
        format="table",
        sheets=sheets,
        sheet=sheet or (sheets[0] if sheets else None),
        header_row=header,
        name_column=guess_name_column(columns, body_rows),
        columns=[
            TableColumn(
                index=c.index,
                header=c.header,
                detected_type=c.detected_type,
                present=c.present,
                numeric_share=round(c.numeric_share, 4),
                percent=c.percent,
                categories=list(c.categories),
                examples=list(c.examples),
            )
            for c in columns
        ],
        rows=kept,
        truncated=len(rows) > limit or any(len(r) > MAX_COLUMNS for r in rows),
    )


@router.post("/build", response_model=Project)
def build_table_project(request: TableBuildRequest) -> Project:
    """Build a project from an inspected table and the user's column choices."""
    specs = [
        ColumnSpec(index=c.index, type=c.type, name=c.name, direction=c.direction, categories=c.categories)
        for c in request.columns
    ]
    return compute(
        lambda: build_project(
            request.rows,
            request.header_row,
            request.name_column,
            specs,
            filename=request.filename,
            sheet=request.sheet,
        )
    )
