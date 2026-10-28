"""Small helpers shared by the spreadsheet readers."""

from __future__ import annotations

import math
from pathlib import Path
from typing import BinaryIO

from ..errors import Location

Rows = list[list[object]]


def column_letter(index: int) -> str:
    """Excel column name for a 1-based index: 1 -> 'A', 27 -> 'AA'."""
    letters = ""
    while index:
        index, rest = divmod(index - 1, 26)
        letters = chr(ord("A") + rest) + letters
    return letters


def cell_ref(row: int, col: int) -> str:
    """Excel-style reference for 0-based indices, e.g. (2, 1) -> 'B3'."""
    return f"{column_letter(col + 1)}{row + 1}"


def cell_location(row: int, col: int) -> Location:
    return Location("cell", ref=cell_ref(row, col))


def as_text(value: object) -> str | None:
    """Cell content as JSON-safe text for error parameters."""
    return None if value is None else str(value)


def is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def parse_number(value: object) -> float | None:
    """Return a finite float, or None if the cell is not a usable number."""
    if isinstance(value, bool) or is_blank(value):
        return None
    try:
        number = float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


Source = str | Path | BinaryIO


def read_xlsx_rows(source: Source, prefer_sheet: str | None = None) -> Rows:
    """Read cell values (formulas resolved to cached results), padded to equal width.

    Uses the sheet named ``prefer_sheet`` if present, else the first sheet.
    """
    # Imported here, not at module level: importing openpyxl is slow in the browser engine (ADR 0009).
    import openpyxl

    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    try:
        sheets = workbook.sheetnames
        worksheet = workbook[prefer_sheet] if prefer_sheet in sheets else workbook.worksheets[0]
        rows: Rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    width = max((len(row) for row in rows), default=0)
    return [row + [None] * (width - len(row)) for row in rows]
