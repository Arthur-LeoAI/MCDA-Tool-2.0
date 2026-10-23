import io

import openpyxl
import pytest
from conftest import LEGACY_DIR

from mcda_core.io import read_problem, write_problem

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def upload(client, filename: str, data: bytes):
    return client.post("/api/v1/problems/import", files={"file": (filename, data)})


def problem_bytes(problem, suffix: str) -> bytes:
    buffer = io.BytesIO()
    write_problem(problem, buffer, suffix=suffix)
    return buffer.getvalue()


@pytest.mark.parametrize("file_format", ["xlsx", "csv"])
def test_template_download(client, file_format):
    response = client.get("/api/v1/template", params={"format": file_format})
    assert response.status_code == 200
    assert f'filename="mcda-template.{file_format}"' in response.headers["content-disposition"]
    problem = read_problem(io.BytesIO(response.content), suffix=f".{file_format}")
    assert problem.criterion_names == ("Cost", "Quality", "Risk")


@pytest.mark.parametrize("suffix", [".xlsx", ".csv"])
def test_import_template(client, cars, cars_json, suffix):
    response = upload(client, f"cars{suffix}", problem_bytes(cars, suffix))
    assert response.status_code == 200
    body = response.json()
    assert body["source_format"] == "template"
    assert body["filename"] == f"cars{suffix}"
    assert body["problem"] == cars_json
    assert body["legacy"] is None


def test_import_legacy_and_reproduce_its_ranking(client):
    path = LEGACY_DIR / "MCDA ELT V7_Equal_G_Top1.xlsx"
    response = upload(client, path.name, path.read_bytes())
    assert response.status_code == 200
    body = response.json()
    assert body["source_format"] == "legacy"
    assert len(body["problem"]["alternatives"]) == 7
    assert len(body["problem"]["criteria"]) == 34
    reference = body["legacy"]["reference_order"]

    evaluated = client.post(
        "/api/v1/evaluate",
        json={
            "problem": body["problem"],
            "methods": [{"method": "weighted_sum", "normalization": "none"}],
        },
    ).json()["results"][0]
    assert evaluated["order"] == reference


def test_import_reports_cell_errors(client):
    data = b"Criterion,c1,c2\nDirection,benefit,sideways\nWeight,1,1\nA,1,n/a\nB,2,3\n"
    response = upload(client, "bad.csv", data)
    assert response.status_code == 422
    found = {(i["code"], i["location"]["ref"]) for i in response.json()["issues"]}
    assert found == {("invalid_direction", "C2"), ("value_not_number", "C4")}


def test_import_rejects_other_file_types(client):
    response = upload(client, "data.json", b"{}")
    assert response.status_code == 415


def test_import_rejects_large_files(client):
    response = upload(client, "big.csv", b"x" * (200 * 1024 + 1))
    assert response.status_code == 413


def test_import_rejects_corrupt_workbook(client):
    response = upload(client, "broken.xlsx", b"not a zip file")
    assert response.status_code == 422
    assert "could not read broken.xlsx" in response.json()["detail"]


def test_import_strips_client_paths(client, cars):
    response = upload(client, "../../secret/cars.csv", problem_bytes(cars, ".csv"))
    assert response.json()["filename"] == "cars.csv"


@pytest.mark.parametrize("file_format", ["xlsx", "csv"])
def test_export_problem_round_trip(client, cars, cars_json, file_format):
    response = client.post("/api/v1/export/problem", params={"format": file_format}, json=cars_json)
    assert response.status_code == 200
    assert read_problem(io.BytesIO(response.content), suffix=f".{file_format}").criteria == cars.criteria


def test_export_report(client, cars_json):
    response = client.post(
        "/api/v1/export/report",
        json={
            "problem": cars_json,
            "methods": [{"method": "weighted_sum"}, {"method": "compromise_programming", "p": 1}],
        },
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Results", "Stability (full order)", "Stability (top 1)", "Pareto", "Decision"]
    header = next(workbook["Results"].iter_rows(values_only=True))
    assert header[1] == "Weighted sum (min_max) score (higher is better)"
    assert header[3] == "CP p=1 (min_max) score (lower is better)"


def test_export_report_without_optional_sheets(client, cars_json):
    response = client.post(
        "/api/v1/export/report",
        json={
            "problem": cars_json,
            "methods": [{"method": "topsis"}],
            "stability": None,
            "include_pareto": False,
        },
    )
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Results", "Decision"]
