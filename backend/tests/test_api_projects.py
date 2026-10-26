import io

import openpyxl
import pytest
from conftest import LEGACY_DIR

from mcda_core.io import write_problem
from mcda_core.project import project_from_problem


@pytest.fixture
def cars_project(cars) -> dict:
    return project_from_problem(cars).model_dump(mode="json")


def test_migrate_v1_problem(client, cars_json):
    response = client.post("/api/v1/projects/migrate", json={**cars_json, "normalization": "min_max"})
    assert response.status_code == 200
    project = response.json()
    assert (project["format"], project["version"]) == ("mcda-project", 2)
    assert project["app_version"]
    assert [c["name"] for c in project["criteria"]] == ["power", "price"]


def test_migrate_rejects_newer_versions_and_garbage(client):
    newer = client.post("/api/v1/projects/migrate", json={"format": "mcda-project", "version": 7})
    assert (newer.status_code, newer.json()["code"]) == (422, "project_version_unsupported")
    garbage = client.post("/api/v1/projects/migrate", json={"hello": "world"})
    assert garbage.json()["code"] == "project_unreadable"


def test_validate(client, cars_project):
    assert client.post("/api/v1/projects/validate", json=cars_project).json() == {"valid": True}
    cars_project["data"]["c1"]["a1"] = "fast"
    response = client.post("/api/v1/projects/validate", json=cars_project)
    assert response.status_code == 422
    (issue,) = response.json()["issues"]
    assert issue["code"] == "value_not_number"
    assert issue["location"] == {"kind": "value", "alternative": "A", "criterion": "power"}


def test_evaluate_returns_ids_values_and_contributions(client, cars_project):
    cars_project["analysis"]["methods"] = [
        {"method": "weighted_sum"},
        {"method": "topsis", "basis": "raw"},
        {"method": "compromise_programming", "p": "inf"},
    ]
    response = client.post("/api/v1/projects/evaluate", json=cars_project)
    assert response.status_code == 200
    body = response.json()
    assert body["eligible"] == ["a1", "a2", "a3"]
    assert body["weights"] == {"c1": 0.5, "c2": 0.5}
    assert body["values"]["a3"] == pytest.approx({"c1": 0.4, "c2": 1.0})
    weighted, topsis, cp = body["results"]
    assert weighted["order"] == ["a3", "a1", "a2"]
    c = next(a for a in weighted["alternatives"] if a["id"] == "a3")
    assert (c["name"], c["rank"]) == ("C", 1)
    assert c["contributions"] == pytest.approx({"c1": 0.2, "c2": 0.5})
    assert topsis["label"] == "TOPSIS (classic)"
    assert cp["higher_is_better"] is False
    assert body["pareto"]["front"] == ["a2", "a3"]


def test_evaluate_reports_exclusions_and_imputations(client, cars_project):
    cars_project["data"]["c1"]["a2"] = None
    cars_project["criteria"][0]["missing"] = "exclude"
    body = client.post("/api/v1/projects/evaluate", json=cars_project).json()
    assert body["eligible"] == ["a1", "a3"]
    (excluded,) = body["excluded"]
    assert excluded["alternative_id"] == "a2"
    assert excluded["reasons"][0]["code"] == "missing_value"

    cars_project["criteria"][0]["missing"] = "worst"
    body = client.post("/api/v1/projects/evaluate", json=cars_project).json()
    assert body["imputed"] == [{"alternative_id": "a2", "criterion_id": "c1", "policy": "worst"}]


def test_groups_without_weights_are_equal(client, cars_project):
    cars_project["groups"] = [{"id": "g1", "name": "Group"}]
    cars_project["criteria"][0]["group_id"] = "g1"
    response = client.post("/api/v1/projects/evaluate", json=cars_project)
    assert response.status_code == 200
    # Upper level: g1 and the ungrouped criteria, equally weighted.
    assert response.json()["weights"] == {"c1": 0.5, "c2": 0.5}


def test_stability(client, cars_project):
    response = client.post("/api/v1/projects/stability", json={"project": cars_project, "scope": "top1"})
    body = response.json()
    assert body["scope"] == "top1"
    assert [i["criterion_id"] for i in body["intervals"]] == ["c1", "c2"]


def test_import_template_and_legacy(client, cars):
    buffer = io.BytesIO()
    write_problem(cars, buffer, suffix=".csv")
    body = client.post("/api/v1/projects/import", files={"file": ("cars.csv", buffer.getvalue())}).json()
    assert body["source_format"] == "template"
    assert body["project"]["criteria"][0]["value_function"]["bounds"] == "data"
    assert body["project"]["source"]["filename"] == "cars.csv"

    path = LEGACY_DIR / "MCDA ELT V7_Equal_G_Top1.xlsx"
    body = client.post("/api/v1/projects/import", files={"file": (path.name, path.read_bytes())}).json()
    project = body["project"]
    assert body["source_format"] == "legacy"
    assert project["criteria"][0]["value_function"]["bounds"] == [0.0, 100.0]
    reference = project["source"]["legacy"]["reference_order"]
    evaluated = client.post("/api/v1/projects/evaluate", json=project).json()
    assert evaluated["results"][0]["order"] == reference


def test_report(client, cars_project):
    response = client.post("/api/v1/projects/report", json={"project": cars_project})
    assert response.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Results", "Stability (full order)", "Stability (top 1)", "Pareto", "Decision"]
    decision = list(workbook["Decision"].iter_rows(values_only=True))
    assert decision[3] == ("A", 100, 30)  # raw data, not value scores


def test_report_leaves_out_decision_sheet_for_non_numeric_data(client, cars_project):
    cars_project["criteria"].append({"id": "c3", "name": "eco", "type": "boolean", "direction": "benefit"})
    cars_project["data"]["c3"] = {"a1": True, "a2": False, "a3": True}
    response = client.post(
        "/api/v1/projects/report", json={"project": cars_project, "include_stability": False, "include_pareto": False}
    )
    assert openpyxl.load_workbook(io.BytesIO(response.content)).sheetnames == ["Results"]
