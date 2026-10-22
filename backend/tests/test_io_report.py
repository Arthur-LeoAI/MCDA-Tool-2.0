import math

import numpy as np
import openpyxl

from mcda_core import (
    StabilityScope,
    pareto_analysis,
    topsis,
    weight_stability_intervals,
    weighted_sum,
)
from mcda_core.io import read_problem, write_report
from mcda_core.io.report import UNBOUNDED


def test_report_contents_and_reimport(cars, tmp_path):
    path = tmp_path / "report.xlsx"
    stability = {scope: weight_stability_intervals(cars, scope) for scope in StabilityScope}
    write_report(
        path,
        cars,
        [("Weighted sum", weighted_sum(cars)), ("TOPSIS", topsis(cars))],
        stability=stability,
        pareto=pareto_analysis(cars),
    )

    workbook = openpyxl.load_workbook(path)
    assert workbook.sheetnames == ["Results", "Stability (full order)", "Stability (top 1)", "Pareto", "Decision"]
    results = list(workbook["Results"].iter_rows(values_only=True))
    assert results[0][:3] == ("Alternative", "Weighted sum score (higher is better)", "Weighted sum rank")
    assert results[3][0] == "C" and results[3][2] == 1

    uppers = [row[3] for row in workbook["Stability (top 1)"].iter_rows(min_row=2, values_only=True)]
    expected = [UNBOUNDED if math.isinf(i.upper) else i.upper for i in stability[StabilityScope.TOP1]]
    assert uppers == expected

    pareto = {row[0]: row[1:] for row in workbook["Pareto"].iter_rows(min_row=2, values_only=True)}
    assert pareto["A"] == ("no", None, "C")

    reimported = read_problem(path)
    assert reimported.alternatives == cars.alternatives
    assert reimported.criteria == cars.criteria
    np.testing.assert_array_equal(reimported.matrix, cars.matrix)
