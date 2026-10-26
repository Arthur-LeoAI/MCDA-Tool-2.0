"""Arbitrary tables: parsing human-written values, guessing structure, building projects, data quality."""

import io

import openpyxl
import pytest

from mcda_core import Direction, Location, McdaError
from mcda_core.project import parse_project, prepare, project_issues
from mcda_core.project.quality import quality_report
from mcda_core.tables import (
    ColumnSpec,
    build_project,
    column_decimal_mark,
    describe_columns,
    guess_header_row,
    guess_name_column,
    is_blank,
    parse_boolean,
    parse_number,
)
from mcda_core.tables.parse import ParsedNumber

# A typical export: a note above the header, money, percentages, a European
# number column, levels, yes/no, placeholders for missing values, an extra column.
MESSY = [
    ["Product research export 2026-10", None, None, None, None, None, None, None],
    ["Product", "Price", "Margin", "Monthly sales", "Weight (kg)", "Competition", "FBA eligible", "Notes"],
    ["Silicone mat", "$12.99", "32%", "1,250", "0,45", "low", "yes", "best seller"],
    ["Bamboo box", "$24.50", "18%", "640", "1,2", "high", "no", ""],
    ["LED strip", "$9.90", "41%", "3,900", "0,3", "medium", "yes", "fragile"],
    ["Pet brush", "-", "27%", "n/a", "0,2", "low", "yes", None],
    ["Desk lamp", "$45.00", "22%", "380", "1.234,5", "high", "no", "heavy"],
]


# --- Parsing ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("text", "value", "percent"),
    [
        ("$1,234.50", 1234.5, False),
        ("15%", 15.0, True),
        ("1.234,5", 1234.5, False),
        ("(300)", -300.0, False),
        ("12,5", 12.5, False),
        ("1,234", 1234.0, False),
        ("1.234.567", 1234567.0, False),
        ("1.234", 1.234, False),
        ("€ 9,99", 9.99, False),
        ("12 USD", 12.0, False),
        ("-7", -7.0, False),
        (" 1" + chr(0xA0) + "234 ", 1234.0, False),
        (3, 3.0, False),
    ],
)
def test_numbers_as_people_write_them(text, value, percent):
    parsed = parse_number(text)
    assert parsed is not None
    assert parsed.value == pytest.approx(value)
    assert parsed.percent is percent


@pytest.mark.parametrize("text", ["abc", "1.2.3,4,5", "--5", "", True, None, "1e309"])
def test_non_numbers(text):
    assert parse_number(text) is None


def test_column_decimal_mark_resolves_ambiguous_values():
    assert column_decimal_mark(["0,45", "1.234,5"]) == ","
    assert parse_number("1.234", ",") == ParsedNumber(1234.0)
    assert parse_number("1,234", ".") == ParsedNumber(1234.0)
    assert column_decimal_mark(["1", "2.5"]) is None


def test_booleans_and_placeholders():
    assert [parse_boolean(v) for v in ["Yes", "nein", "是", "否", 1, 0, True, "maybe"]] == [
        True,
        False,
        True,
        False,
        True,
        False,
        True,
        None,
    ]
    assert all(is_blank(v) for v in [None, "", "  ", "-", "—", "n/a", "NA"])
    assert not is_blank("0")


# --- Structure ---------------------------------------------------------------------------


def test_detects_header_row_types_and_name_column():
    assert guess_header_row(MESSY) == 1
    columns = describe_columns(MESSY, 1)
    types = {c.header: c.detected_type for c in columns}
    assert types == {
        "Product": "categorical",
        "Price": "numeric",
        "Margin": "numeric",
        "Monthly sales": "numeric",
        "Weight (kg)": "numeric",
        "Competition": "ordinal",
        "FBA eligible": "boolean",
        "Notes": "categorical",
    }
    competition = next(c for c in columns if c.header == "Competition")
    assert competition.categories == ("low", "medium", "high")  # vocabulary order, not appearance
    assert next(c for c in columns if c.header == "Margin").percent
    assert guess_name_column(columns, rows=5) == 0


def test_header_defaults_to_first_row():
    assert guess_header_row([["a", "b"], ["x", "y"]]) == 0


# --- Building a project --------------------------------------------------------------------


def messy_project():
    return build_project(
        MESSY,
        header_row=1,
        name_column=0,
        columns=[
            ColumnSpec(1, "numeric", direction=Direction.COST),
            ColumnSpec(2, "numeric"),
            ColumnSpec(3, "numeric", name="Sales / month"),
            ColumnSpec(4, "numeric", direction=Direction.COST),
            ColumnSpec(5, "ordinal", categories=["high", "medium", "low"]),  # low competition is best
            ColumnSpec(6, "boolean"),
        ],
        filename="export.xlsx",
        sheet="Sheet1",
    )


def test_build_reads_values_and_settings():
    project = messy_project()
    assert [a.name for a in project.alternatives] == [
        "Silicone mat",
        "Bamboo box",
        "LED strip",
        "Pet brush",
        "Desk lamp",
    ]
    names = [c.name for c in project.criteria]
    assert names == ["Price", "Margin", "Sales / month", "Weight (kg)", "Competition", "FBA eligible"]
    price, margin, sales, weight, competition, fba = project.criteria
    assert project.data[price.id] == {"a1": 12.99, "a2": 24.5, "a3": 9.9, "a4": None, "a5": 45.0}
    assert margin.unit == "%" and project.raw(margin.id, "a1") == 32.0
    assert project.raw(sales.id, "a3") == 3900.0 and project.raw(sales.id, "a4") is None
    assert project.data[weight.id]["a5"] == 1234.5 and project.data[weight.id]["a1"] == 0.45  # European column
    assert [(c.label, c.score) for c in competition.categories or []] == [("high", 0.0), ("medium", 0.5), ("low", 1.0)]
    assert project.raw(fba.id, "a2") is False
    assert project.meta.name == "export"
    assert project.source is not None and project.source.kind == "table" and project.source.header_row == 2
    assert project_issues(project) == []
    prepare(project)  # analysable as built


def test_unreadable_cells_are_kept_for_validation():
    rows = [["name", "x"], ["A", "12"], ["B", "twelve"]]
    project = build_project(rows, 0, 0, [ColumnSpec(1, "numeric")])
    (issue,) = project_issues(project)
    assert issue.code == "value_not_number"
    assert issue.location == Location("value", alternative="B", criterion="x")


def test_build_checks_column_choices():
    with pytest.raises(McdaError, match="at least one column"):
        build_project(MESSY, 1, 0, [])
    with pytest.raises(McdaError) as excinfo:
        build_project(MESSY, 1, 0, [ColumnSpec(42, "numeric")])
    assert excinfo.value.params == {"index": 42}


# --- Data quality --------------------------------------------------------------------------


def test_quality_report_findings():
    header: list[object] = ["name", "sales", "flat", "size"]
    rows: list[list[object]] = [header] + [
        [f"P{i}", value, 5, size]  # type: list[object]
        for i, (value, size) in enumerate([(100, "a"), (120, "b"), (110, "a"), (90, "b"), (20000, "a")])
    ]
    rows.append(["P5", None, 5, "a"])
    project = build_project(
        rows, 0, 0, [ColumnSpec(1, "numeric"), ColumnSpec(2, "numeric"), ColumnSpec(3, "categorical")]
    )
    found = {(f.code, f.location.name if f.location else None): f.params for f in quality_report(project)}
    assert found[("quality_missing", "sales")] == {"count": 1, "policy": "worst"}
    assert found[("quality_outliers", "sales")] == {"count": 1, "alternatives": "P4"}
    assert found[("quality_suggest_log", "sales")] == {"ratio": 222}
    assert ("quality_constant", "flat") in found
    assert ("quality_categories_equal", "size") in found


# --- API -----------------------------------------------------------------------------------


def messy_xlsx() -> bytes:
    workbook = openpyxl.Workbook()
    notes = workbook.worksheets[0]
    notes.title = "Notes"
    notes["A1"] = "nothing here"
    sheet = workbook.create_sheet("Products")
    for row in MESSY:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_api_inspect_build_quality(client):
    upload = {"file": ("export.xlsx", messy_xlsx())}
    first = client.post("/api/v1/tables/inspect", files=upload).json()
    assert first["format"] == "table"
    assert first["sheets"] == ["Notes", "Products"] and first["sheet"] == "Notes"

    body = client.post("/api/v1/tables/inspect", params={"sheet": "Products"}, files=upload).json()
    assert (body["sheet"], body["header_row"], body["name_column"]) == ("Products", 1, 0)
    assert [c["detected_type"] for c in body["columns"]][1:7] == [
        "numeric",
        "numeric",
        "numeric",
        "numeric",
        "ordinal",
        "boolean",
    ]
    assert body["rows"][2][1] == "$12.99"

    request = {
        "rows": body["rows"],
        "header_row": 1,
        "name_column": 0,
        "columns": [
            {"index": 1, "type": "numeric", "direction": "cost"},
            {"index": 5, "type": "ordinal", "categories": ["high", "medium", "low"]},
        ],
        "filename": "export.xlsx",
        "sheet": "Products",
    }
    project = client.post("/api/v1/tables/build", json=request).json()
    assert [c["name"] for c in project["criteria"]] == ["Price", "Competition"]
    assert parse_project(project).raw("c1", "a1") == 12.99

    quality = client.post("/api/v1/projects/quality", json=project).json()
    assert {(q["code"], q["location"]["name"]) for q in quality} == {("quality_missing", "Price")}


def test_api_inspect_recognises_templates_and_legacy(client, cars):
    from conftest import LEGACY_DIR

    from mcda_core.io import write_problem

    buffer = io.BytesIO()
    write_problem(cars, buffer, suffix=".csv")
    body = client.post("/api/v1/tables/inspect", files={"file": ("cars.csv", buffer.getvalue())}).json()
    assert body["format"] == "template" and body["project"]["criteria"][0]["name"] == "power"

    path = LEGACY_DIR / "MCDA ELT V7_Equal_G_Top1.xlsx"
    body = client.post("/api/v1/tables/inspect", files={"file": (path.name, path.read_bytes())}).json()
    assert body["format"] == "legacy" and len(body["project"]["criteria"]) == 34


def test_api_inspect_errors(client):
    upload = {"file": ("export.xlsx", messy_xlsx())}
    missing = client.post("/api/v1/tables/inspect", params={"sheet": "Nope"}, files=upload)
    assert (missing.status_code, missing.json()["code"]) == (422, "sheet_not_found")
    empty = client.post("/api/v1/tables/inspect", files={"file": ("empty.csv", b"\n\n")})
    assert empty.json()["code"] == "table_empty"
